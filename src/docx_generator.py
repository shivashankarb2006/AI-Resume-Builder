"""
AI Resume Builder - Professional ATS DOCX Document Generator
Builds beautifully styled, ATS-compliant Microsoft Word (.docx) documents using python-docx.
"""

import io
from typing import Dict, Any
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls


# Professional Palette
COLOR_PRIMARY = RGBColor(15, 23, 42)      # Deep Navy / Slate 900
COLOR_SECONDARY = RGBColor(51, 65, 85)   # Slate 700
COLOR_MUTED = RGBColor(71, 85, 105)      # Slate 600
COLOR_BODY = RGBColor(17, 24, 39)        # Neutral Dark


def add_horizontal_rule(paragraph):
    """Add a clean bottom border line under section headings for ATS readability."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                     r'<w:bottom w:val="single" w:sz="6" w:space="2" w:color="CBD5E1"/>'
                     r'</w:pBdr>')
    pPr.append(pBdr)


def set_run_font(run, name="Arial", size_pt=10.5, color=COLOR_BODY, bold=False, italic=False):
    """Standardize font properties across runs."""
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic


def add_section_header(doc: Document, title: str):
    """Add a standardized uppercase section heading with an ATS-friendly divider line."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title.upper())
    set_run_font(run, name="Arial", size_pt=11, color=COLOR_PRIMARY, bold=True)
    add_horizontal_rule(p)
    return p


def build_docx_resume(resume: Dict[str, Any]) -> io.BytesIO:
    """
    Generate an ATS-optimized .docx document from structured resume data.
    Returns an in-memory BytesIO buffer ready for st.download_button.
    """
    doc = Document()

    # Set 0.6 inch margins (standard ATS recommendation)
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    p_info = resume.get("personal_info", {})

    # 1. HEADER - NAME
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_p.paragraph_format.space_before = Pt(0)
    name_p.paragraph_format.space_after = Pt(2)
    name_run = name_p.add_run(p_info.get("full_name", "Resume").upper())
    set_run_font(name_run, name="Arial", size_pt=18, color=COLOR_PRIMARY, bold=True)

    # 2. HEADER - TITLE
    if p_info.get("professional_title"):
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_p.paragraph_format.space_before = Pt(0)
        title_p.paragraph_format.space_after = Pt(4)
        title_run = title_p.add_run(p_info["professional_title"])
        set_run_font(title_run, name="Arial", size_pt=11, color=COLOR_SECONDARY, bold=True)

    # 3. HEADER - CONTACT INFO LINE
    contact_parts = [
        p_info.get("email"),
        p_info.get("phone"),
        p_info.get("location"),
        p_info.get("linkedin"),
        p_info.get("github"),
        p_info.get("portfolio"),
    ]
    contact_str = " | ".join([c for c in contact_parts if c])
    if contact_str:
        contact_p = doc.add_paragraph()
        contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact_p.paragraph_format.space_before = Pt(0)
        contact_p.paragraph_format.space_after = Pt(8)
        contact_run = contact_p.add_run(contact_str)
        set_run_font(contact_run, name="Arial", size_pt=9.5, color=COLOR_MUTED)

    # 4. PROFESSIONAL SUMMARY
    summary = resume.get("professional_summary")
    if summary and summary.strip():
        add_section_header(doc, "Professional Summary")
        p_sum = doc.add_paragraph()
        p_sum.paragraph_format.space_before = Pt(2)
        p_sum.paragraph_format.space_after = Pt(6)
        p_sum.paragraph_format.line_spacing = 1.15
        run_sum = p_sum.add_run(summary.strip())
        set_run_font(run_sum, name="Arial", size_pt=10, color=COLOR_BODY)

    # 5. TECHNICAL SKILLS
    skills = resume.get("skills", {})
    if skills and any(skills.values()):
        add_section_header(doc, "Technical Skills")
        for category, items in skills.items():
            cat_name = category.replace("_", " ").title()
            if isinstance(items, list) and items:
                items_str = ", ".join(items)
            elif isinstance(items, str) and items.strip():
                items_str = items.strip()
            else:
                continue

            p_sk = doc.add_paragraph()
            p_sk.paragraph_format.space_before = Pt(1)
            p_sk.paragraph_format.space_after = Pt(2)
            p_sk.paragraph_format.line_spacing = 1.15

            run_cat = p_sk.add_run(f"{cat_name}: ")
            set_run_font(run_cat, name="Arial", size_pt=10, color=COLOR_PRIMARY, bold=True)

            run_items = p_sk.add_run(items_str)
            set_run_font(run_items, name="Arial", size_pt=10, color=COLOR_BODY)

    # 6. WORK EXPERIENCE
    experience = resume.get("experience", [])
    if experience:
        add_section_header(doc, "Work Experience")
        for exp in experience:
            if not exp.get("role") and not exp.get("company"):
                continue

            # Header Line: Role | Company (Duration)
            p_exp = doc.add_paragraph()
            p_exp.paragraph_format.space_before = Pt(6)
            p_exp.paragraph_format.space_after = Pt(2)
            p_exp.paragraph_format.keep_with_next = True

            run_role = p_exp.add_run(exp.get("role", "Software Engineer"))
            set_run_font(run_role, name="Arial", size_pt=10.5, color=COLOR_PRIMARY, bold=True)

            run_sep = p_exp.add_run(" | ")
            set_run_font(run_sep, name="Arial", size_pt=10.5, color=COLOR_MUTED)

            run_comp = p_exp.add_run(exp.get("company", "Company"))
            set_run_font(run_comp, name="Arial", size_pt=10.5, color=COLOR_SECONDARY, bold=True)

            meta_parts = []
            if exp.get("location"):
                meta_parts.append(exp["location"])
            if exp.get("duration"):
                meta_parts.append(exp["duration"])

            if meta_parts:
                run_dur = p_exp.add_run(f" ({' - '.join(meta_parts)})")
                set_run_font(run_dur, name="Arial", size_pt=9.5, color=COLOR_MUTED, italic=True)

            # Bullet points
            for bp in exp.get("bullet_points", []):
                if bp and bp.strip():
                    p_bp = doc.add_paragraph(style="List Bullet")
                    p_bp.paragraph_format.space_before = Pt(1)
                    p_bp.paragraph_format.space_after = Pt(2)
                    p_bp.paragraph_format.line_spacing = 1.15
                    run_bp = p_bp.add_run(bp.strip())
                    set_run_font(run_bp, name="Arial", size_pt=10, color=COLOR_BODY)

    # 7. TECHNICAL PROJECTS
    projects = resume.get("projects", [])
    if projects:
        add_section_header(doc, "Technical Projects")
        for proj in projects:
            if not proj.get("name"):
                continue

            p_pr = doc.add_paragraph()
            p_pr.paragraph_format.space_before = Pt(6)
            p_pr.paragraph_format.space_after = Pt(2)
            p_pr.paragraph_format.keep_with_next = True

            run_pname = p_pr.add_run(proj.get("name", "Project"))
            set_run_font(run_pname, name="Arial", size_pt=10.5, color=COLOR_PRIMARY, bold=True)

            if proj.get("technologies"):
                run_tech = p_pr.add_run(f" | {proj['technologies']}")
                set_run_font(run_tech, name="Arial", size_pt=9.5, color=COLOR_MUTED, italic=True)

            if proj.get("urls"):
                run_url = p_pr.add_run(f" ({proj['urls']})")
                set_run_font(run_url, name="Arial", size_pt=9, color=COLOR_MUTED)

            for bp in proj.get("bullet_points", []):
                if bp and bp.strip():
                    p_bp = doc.add_paragraph(style="List Bullet")
                    p_bp.paragraph_format.space_before = Pt(1)
                    p_bp.paragraph_format.space_after = Pt(2)
                    p_bp.paragraph_format.line_spacing = 1.15
                    run_bp = p_bp.add_run(bp.strip())
                    set_run_font(run_bp, name="Arial", size_pt=10, color=COLOR_BODY)

    # 8. EDUCATION
    education = resume.get("education", [])
    if education:
        add_section_header(doc, "Education")
        for edu in education:
            if not edu.get("institution") and not edu.get("degree"):
                continue

            p_ed = doc.add_paragraph()
            p_ed.paragraph_format.space_before = Pt(4)
            p_ed.paragraph_format.space_after = Pt(2)

            deg_title = f"{edu.get('degree', '')}"
            if edu.get("field_of_study"):
                deg_title += f" in {edu['field_of_study']}"

            run_deg = p_ed.add_run(deg_title)
            set_run_font(run_deg, name="Arial", size_pt=10, color=COLOR_PRIMARY, bold=True)

            run_inst = p_ed.add_run(f" - {edu.get('institution', '')}")
            set_run_font(run_inst, name="Arial", size_pt=10, color=COLOR_SECONDARY)

            if edu.get("year"):
                run_yr = p_ed.add_run(f" ({edu['year']})")
                set_run_font(run_yr, name="Arial", size_pt=9.5, color=COLOR_MUTED, italic=True)

            if edu.get("gpa"):
                run_gpa = p_ed.add_run(f" | GPA: {edu['gpa']}")
                set_run_font(run_gpa, name="Arial", size_pt=9.5, color=COLOR_MUTED)

            if edu.get("highlights"):
                p_hi = doc.add_paragraph(style="List Bullet")
                p_hi.paragraph_format.space_before = Pt(1)
                p_hi.paragraph_format.space_after = Pt(2)
                run_hi = p_hi.add_run(edu["highlights"])
                set_run_font(run_hi, name="Arial", size_pt=9.5, color=COLOR_BODY)

    # 9. CERTIFICATIONS
    certs = [c for c in resume.get("certifications", []) if c.get("name")]
    if certs:
        add_section_header(doc, "Certifications")
        for cert in certs:
            p_c = doc.add_paragraph(style="List Bullet")
            p_c.paragraph_format.space_before = Pt(1)
            p_c.paragraph_format.space_after = Pt(2)

            run_cnm = p_c.add_run(cert["name"])
            set_run_font(run_cnm, name="Arial", size_pt=10, color=COLOR_PRIMARY, bold=True)

            org_str = f" - {cert.get('organization', '')}" if cert.get("organization") else ""
            dt_str = f" ({cert.get('date', '')})" if cert.get("date") else ""
            run_corg = p_c.add_run(f"{org_str}{dt_str}")
            set_run_font(run_corg, name="Arial", size_pt=9.5, color=COLOR_BODY)

    # 10. ACHIEVEMENTS
    achs = [a for a in resume.get("achievements", []) if a.strip()]
    if achs:
        add_section_header(doc, "Achievements & Honors")
        for ach in achs:
            p_a = doc.add_paragraph(style="List Bullet")
            p_a.paragraph_format.space_before = Pt(1)
            p_a.paragraph_format.space_after = Pt(2)
            run_a = p_a.add_run(ach.strip())
            set_run_font(run_a, name="Arial", size_pt=10, color=COLOR_BODY)

    # Save to buffer
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
