"""
python-docx Generator for AI Resume Builder.
Produces fully editable, ATS-friendly Word (.docx) documents with correct margins,
consistent typography, bullet points, and clickable hyperlinks.
"""
import io
from typing import Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from models.resume_schema import ResumeData
from .templates import TEMPLATE_CONFIGS, TEMPLATE_CLASSIC_ATS


def _hex_to_rgb(hex_str: str) -> RGBColor:
    """Convert hex string (e.g. '#2563eb' or '2563eb') to RGBColor."""
    hex_str = hex_str.lstrip("#")
    if len(hex_str) != 6:
        return RGBColor(0, 0, 0)
    r = int(hex_str[0:2], 16)
    g = int(hex_str[2:4], 16)
    b = int(hex_str[4:6], 16)
    return RGBColor(r, g, b)


def add_hyperlink(paragraph, url: str, text: str, color_hex: str = "2563eb"):
    """
    Add a working clickable hyperlink into a python-docx paragraph.
    Uses native Word OpenXML relationship elements.
    """
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    # Set link color
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color_hex.lstrip("#"))
    rPr.append(c)

    # Set underline
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)

    new_run.append(rPr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)

    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _add_section_header(doc: Document, title: str, font_name: str, heading_rgb: RGBColor):
    """Add standardized ATS section heading with bottom border formatting."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True

    run = p.add_run(title.upper())
    run.font.name = font_name
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = heading_rgb

    # Add bottom border to heading paragraph
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")  # 0.75 pt
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "94A3B8")
    pBdr.append(bottom)
    pPr.append(pBdr)


def generate_resume_docx(resume_data: ResumeData) -> bytes:
    """
    Generate an editable, ATS-optimized Microsoft Word (.docx) document.
    """
    doc = Document()

    # Template config
    template_name = resume_data.selected_template or TEMPLATE_CLASSIC_ATS
    config = TEMPLATE_CONFIGS.get(template_name, TEMPLATE_CONFIGS[TEMPLATE_CLASSIC_ATS])

    font_name = config["font_family_docx"]
    heading_rgb = _hex_to_rgb(config["heading_color"])
    primary_rgb = _hex_to_rgb(config["primary_color"])
    accent_hex = config["accent_color"]
    is_center = config["header_alignment"] == "center"

    # Set Page Margins to standard 0.5 inch (36 pt)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.page_width = Inches(8.27)   # A4 Width
        section.page_height = Inches(11.69) # A4 Height

    # Set base Normal style font
    style_normal = doc.styles["Normal"]
    style_normal.font.name = font_name
    style_normal.font.size = Pt(9.5)
    style_normal.font.color.rgb = primary_rgb

    # 1. Candidate Name Header
    p_name = doc.add_paragraph()
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(2)
    if is_center:
        p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER

    name_run = p_name.add_run(resume_data.personal_info.full_name or "Your Name")
    name_run.font.name = font_name
    name_run.font.size = Pt(20)
    name_run.font.bold = True
    name_run.font.color.rgb = heading_rgb

    # 2. Contact Info Bar
    p = resume_data.personal_info
    p_contact = doc.add_paragraph()
    p_contact.paragraph_format.space_before = Pt(0)
    p_contact.paragraph_format.space_after = Pt(8)
    if is_center:
        p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER

    contact_items = []
    if p.email:
        contact_items.append(("link", p.email, f"mailto:{p.email}"))
    if p.phone:
        contact_items.append(("text", p.phone, ""))
    if p.location:
        contact_items.append(("text", p.location, ""))
    if p.linkedin_url:
        contact_items.append(("link", "LinkedIn", p.linkedin_url))
    if p.github_url:
        contact_items.append(("link", "GitHub", p.github_url))
    if p.portfolio_url:
        contact_items.append(("link", "Portfolio", p.portfolio_url))

    for idx, (kind, label, target) in enumerate(contact_items):
        if kind == "link":
            add_hyperlink(p_contact, target, label, accent_hex)
        else:
            r = p_contact.add_run(label)
            r.font.name = font_name
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(100, 116, 139)

        if idx < len(contact_items) - 1:
            sep_run = p_contact.add_run("  |  ")
            sep_run.font.name = font_name
            sep_run.font.size = Pt(9.5)
            sep_run.font.color.rgb = RGBColor(148, 163, 184)

    enabled = resume_data.enabled_sections
    section_order = config.get("section_order", [
        "Professional Summary",
        "Education",
        "Skills",
        "Experience",
        "Projects",
        "Certifications",
        "Achievements",
    ])

    for sec in section_order:
        if sec not in enabled:
            continue

        if sec == "Professional Summary" and resume_data.summary.strip():
            _add_section_header(doc, "Professional Summary", font_name, heading_rgb)
            p_sum = doc.add_paragraph()
            p_sum.paragraph_format.space_before = Pt(2)
            p_sum.paragraph_format.space_after = Pt(4)
            p_sum.paragraph_format.line_spacing = 1.15
            r = p_sum.add_run(resume_data.summary.strip())
            r.font.name = font_name
            r.font.size = Pt(9.5)
            r.font.color.rgb = primary_rgb

        elif sec == "Education" and resume_data.education:
            valid_edus = [e for e in resume_data.education if not e.is_empty()]
            if valid_edus:
                _add_section_header(doc, "Education", font_name, heading_rgb)
                for e in valid_edus:
                    p_edu = doc.add_paragraph()
                    p_edu.paragraph_format.space_before = Pt(2)
                    p_edu.paragraph_format.space_after = Pt(1)

                    r_inst = p_edu.add_run(e.institution)
                    r_inst.font.name = font_name
                    r_inst.font.bold = True
                    r_inst.font.size = Pt(10)
                    r_inst.font.color.rgb = heading_rgb

                    dates = f" ({e.start_year} – {e.end_year})" if e.start_year or e.end_year else ""
                    if dates:
                        r_dt = p_edu.add_run(dates)
                        r_dt.font.name = font_name
                        r_dt.font.size = Pt(9)
                        r_dt.font.color.rgb = RGBColor(100, 116, 139)

                    # Degree line
                    p_deg = doc.add_paragraph()
                    p_deg.paragraph_format.space_before = Pt(0)
                    p_deg.paragraph_format.space_after = Pt(2)
                    deg_str = f"{e.degree} in {e.field_of_study}" if e.field_of_study else e.degree
                    if e.gpa:
                        deg_str += f" | GPA: {e.gpa}"
                    r_deg = p_deg.add_run(deg_str)
                    r_deg.font.name = font_name
                    r_deg.font.size = Pt(9.5)

                    if e.relevant_coursework:
                        p_cw = doc.add_paragraph(style="List Bullet")
                        p_cw.paragraph_format.space_before = Pt(0)
                        p_cw.paragraph_format.space_after = Pt(2)
                        p_cw.paragraph_format.line_spacing = 1.1
                        r_lbl = p_cw.add_run("Relevant Coursework: ")
                        r_lbl.font.bold = True
                        r_cw = p_cw.add_run(e.relevant_coursework)
                        r_cw.font.name = font_name
                        r_cw.font.size = Pt(9)

        elif sec == "Skills" and not resume_data.skills.is_empty():
            s = resume_data.skills
            _add_section_header(doc, "Technical Skills", font_name, heading_rgb)

            def _add_skill_line(label: str, items: list):
                if items:
                    p_sk = doc.add_paragraph()
                    p_sk.paragraph_format.space_before = Pt(1)
                    p_sk.paragraph_format.space_after = Pt(1)
                    p_sk.paragraph_format.line_spacing = 1.15
                    r_lbl = p_sk.add_run(f"{label}: ")
                    r_lbl.font.name = font_name
                    r_lbl.font.bold = True
                    r_lbl.font.size = Pt(9.5)
                    r_val = p_sk.add_run(", ".join(items))
                    r_val.font.name = font_name
                    r_val.font.size = Pt(9.5)

            _add_skill_line("Languages", s.programming_languages)
            _add_skill_line("Frameworks & Libraries", s.frameworks)
            _add_skill_line("Developer Tools", s.tools)
            _add_skill_line("AI & Machine Learning", s.ai_ml)
            _add_skill_line("Databases", s.databases)
            for c in s.custom_categories:
                if c.category_name.strip() and c.skills:
                    _add_skill_line(c.category_name, c.skills)

        elif sec == "Experience" and resume_data.experience:
            valid_exps = [e for e in resume_data.experience if not e.is_empty()]
            if valid_exps:
                _add_section_header(doc, "Professional Experience", font_name, heading_rgb)
                for exp in valid_exps:
                    p_exp = doc.add_paragraph()
                    p_exp.paragraph_format.space_before = Pt(3)
                    p_exp.paragraph_format.space_after = Pt(1)
                    p_exp.paragraph_format.keep_with_next = True

                    r_comp = p_exp.add_run(exp.company)
                    r_comp.font.name = font_name
                    r_comp.font.bold = True
                    r_comp.font.size = Pt(10)
                    r_comp.font.color.rgb = heading_rgb

                    if exp.location:
                        r_loc = p_exp.add_run(f" — {exp.location}")
                        r_loc.font.name = font_name
                        r_loc.font.size = Pt(9.5)
                        r_loc.font.color.rgb = RGBColor(100, 116, 139)

                    end_d = "Present" if exp.is_current else exp.end_date
                    dates = f" | ({exp.start_date} – {end_d})"
                    r_d = p_exp.add_run(dates)
                    r_d.font.name = font_name
                    r_d.font.size = Pt(9)
                    r_d.font.color.rgb = RGBColor(100, 116, 139)

                    # Role
                    p_role = doc.add_paragraph()
                    p_role.paragraph_format.space_before = Pt(0)
                    p_role.paragraph_format.space_after = Pt(2)
                    r_role = p_role.add_run(exp.role)
                    r_role.font.name = font_name
                    r_role.font.italic = True
                    r_role.font.size = Pt(9.5)

                    bullets = exp.improved_bullets
                    if not bullets and exp.responsibilities:
                        bullets = [b.strip() for b in exp.responsibilities.split("\n") if b.strip()]

                    for b in bullets:
                        p_b = doc.add_paragraph(style="List Bullet")
                        p_b.paragraph_format.space_before = Pt(0)
                        p_b.paragraph_format.space_after = Pt(2)
                        p_b.paragraph_format.line_spacing = 1.15
                        r_b = p_b.add_run(b)
                        r_b.font.name = font_name
                        r_b.font.size = Pt(9.5)

        elif sec == "Projects" and resume_data.projects:
            valid_projs = [p for p in resume_data.projects if not p.is_empty()]
            if valid_projs:
                _add_section_header(doc, "Technical Projects", font_name, heading_rgb)
                for proj in valid_projs:
                    p_proj = doc.add_paragraph()
                    p_proj.paragraph_format.space_before = Pt(3)
                    p_proj.paragraph_format.space_after = Pt(1)
                    p_proj.paragraph_format.keep_with_next = True

                    r_pname = p_proj.add_run(proj.name)
                    r_pname.font.name = font_name
                    r_pname.font.bold = True
                    r_pname.font.size = Pt(10)
                    r_pname.font.color.rgb = heading_rgb

                    if proj.technologies:
                        r_tech = p_proj.add_run(f" | {', '.join(proj.technologies)}")
                        r_tech.font.name = font_name
                        r_tech.font.italic = True
                        r_tech.font.size = Pt(9)
                        r_tech.font.color.rgb = RGBColor(71, 85, 105)

                    if proj.github_url:
                        p_proj.add_run("  [")
                        add_hyperlink(p_proj, proj.github_url, "GitHub", accent_hex)
                        p_proj.add_run("]")
                    if proj.demo_url:
                        p_proj.add_run("  [")
                        add_hyperlink(p_proj, proj.demo_url, "Demo", accent_hex)
                        p_proj.add_run("]")

                    bullets = proj.improved_bullets
                    if not bullets and proj.key_contributions:
                        bullets = [b.strip() for b in proj.key_contributions.split("\n") if b.strip()]
                    elif not bullets and proj.description:
                        bullets = [proj.description]

                    for b in bullets:
                        p_b = doc.add_paragraph(style="List Bullet")
                        p_b.paragraph_format.space_before = Pt(0)
                        p_b.paragraph_format.space_after = Pt(2)
                        p_b.paragraph_format.line_spacing = 1.15
                        r_b = p_b.add_run(b)
                        r_b.font.name = font_name
                        r_b.font.size = Pt(9.5)

        elif sec == "Certifications" and resume_data.certifications:
            valid_certs = [c for c in resume_data.certifications if not c.is_empty()]
            if valid_certs:
                _add_section_header(doc, "Certifications", font_name, heading_rgb)
                for c in valid_certs:
                    p_c = doc.add_paragraph(style="List Bullet")
                    p_c.paragraph_format.space_before = Pt(0)
                    p_c.paragraph_format.space_after = Pt(2)
                    r_cn = p_c.add_run(c.name)
                    r_cn.font.name = font_name
                    r_cn.font.bold = True
                    r_cn.font.size = Pt(9.5)

                    issuer_dt = ""
                    if c.issuer:
                        issuer_dt += f" — {c.issuer}"
                    if c.issue_date:
                        issuer_dt += f" ({c.issue_date})"
                    if issuer_dt:
                        r_id = p_c.add_run(issuer_dt)
                        r_id.font.name = font_name
                        r_id.font.size = Pt(9.5)

                    if c.credential_url:
                        p_c.add_run("  [")
                        add_hyperlink(p_c, c.credential_url, "Verify", accent_hex)
                        p_c.add_run("]")

        elif sec == "Achievements" and resume_data.achievements:
            valid_achs = [a for a in resume_data.achievements if not a.is_empty()]
            if valid_achs:
                _add_section_header(doc, "Honors & Achievements", font_name, heading_rgb)
                for a in valid_achs:
                    p_a = doc.add_paragraph(style="List Bullet")
                    p_a.paragraph_format.space_before = Pt(0)
                    p_a.paragraph_format.space_after = Pt(2)

                    r_t = p_a.add_run(a.title)
                    r_t.font.name = font_name
                    r_t.font.bold = True
                    r_t.font.size = Pt(9.5)

                    extra = ""
                    if a.issuer_or_event:
                        extra += f" | {a.issuer_or_event}"
                    if a.date:
                        extra += f" ({a.date})"
                    if a.description:
                        extra += f": {a.description}"

                    if extra:
                        r_ex = p_a.add_run(extra)
                        r_ex.font.name = font_name
                        r_ex.font.size = Pt(9.5)

        elif sec == "Volunteer Experience" and resume_data.volunteer:
            valid_vols = [v for v in resume_data.volunteer if not v.is_empty()]
            if valid_vols:
                _add_section_header(doc, "Volunteer & Leadership", font_name, heading_rgb)
                for v in valid_vols:
                    p_v = doc.add_paragraph(style="List Bullet")
                    p_v.paragraph_format.space_before = Pt(0)
                    p_v.paragraph_format.space_after = Pt(2)

                    r_vr = p_v.add_run(f"{v.role} at {v.organization}")
                    r_vr.font.name = font_name
                    r_vr.font.bold = True
                    r_vr.font.size = Pt(9.5)

                    details = ""
                    if v.dates:
                        details += f" ({v.dates})"
                    if v.details:
                        details += f": {v.details}"

                    if details:
                        r_vd = p_v.add_run(details)
                        r_vd.font.name = font_name
                        r_vd.font.size = Pt(9.5)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
