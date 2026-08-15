from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SECTION_NAMES = {
    "PROFESSIONAL SUMMARY",
    "SUMMARY",
    "CAREER OBJECTIVE",
    "EDUCATION",
    "TECHNICAL SKILLS",
    "SKILLS",
    "PROJECTS",
    "EXPERIENCE",
    "WORK EXPERIENCE",
    "INTERNSHIPS",
    "CERTIFICATIONS",
    "ACHIEVEMENTS",
    "LANGUAGES",
}


def set_cell_shading(cell, fill):
    """Set background color of a table cell."""

    tc_pr = cell._tc.get_or_add_tcPr()

    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)

    tc_pr.append(shd)


def add_bottom_border(paragraph, color="000000", size="6"):
    """Add a line below a paragraph."""

    p = paragraph._p
    pPr = p.get_or_add_pPr()

    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")

    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)

    pBdr.append(bottom)
    pPr.append(pBdr)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    """Set table cell margins."""

    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()

    tcMar = tcPr.first_child_found_in("w:tcMar")

    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)

    for margin, value in [
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ]:

        node = tcMar.find(qn(f"w:{margin}"))

        if node is None:
            node = OxmlElement(f"w:{margin}")
            tcMar.append(node)

        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def create_docx(resume_text, template="Classic ATS"):

    document = Document()

    section = document.sections[0]

    # =====================================================
    # TEMPLATE SETTINGS
    # =====================================================

    if template == "Modern Professional":

        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)

        font_name = "Calibri"
        body_size = 10
        heading_size = 11
        name_size = 23

        accent_color = RGBColor(31, 78, 121)

    elif template == "Minimal Student":

        section.top_margin = Inches(0.45)
        section.bottom_margin = Inches(0.45)
        section.left_margin = Inches(0.55)
        section.right_margin = Inches(0.55)

        font_name = "Arial"
        body_size = 9
        heading_size = 10
        name_size = 17

        accent_color = RGBColor(0, 0, 0)

    else:

        # Classic ATS

        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

        font_name = "Arial"
        body_size = 10
        heading_size = 11
        name_size = 20

        accent_color = RGBColor(0, 0, 0)

    # =====================================================
    # DEFAULT STYLE
    # =====================================================

    normal_style = document.styles["Normal"]

    normal_style.font.name = font_name
    normal_style.font.size = Pt(body_size)

    # =====================================================
    # RESUME LINES
    # =====================================================

    lines = resume_text.split("\n")

    if not lines:
        return None

    # =====================================================
    # HEADER
    # =====================================================

    name = lines[0].strip()

    if name:

        if template == "Modern Professional":

            # Modern colored header box

            table = document.add_table(
                rows=1,
                cols=1
            )

            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True

            cell = table.cell(0, 0)

            set_cell_shading(
                cell,
                "1F4E79"
            )

            set_cell_margins(
                cell,
                top=180,
                bottom=180,
                start=180,
                end=180
            )

            cell.vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.CENTER
            )

            paragraph = cell.paragraphs[0]

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT
            )

            run = paragraph.add_run(name)

            run.bold = True
            run.font.name = "Calibri"
            run.font.size = Pt(name_size)
            run.font.color.rgb = RGBColor(
                255,
                255,
                255
            )

            document.add_paragraph()

        elif template == "Minimal Student":

            paragraph = document.add_paragraph()

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            run = paragraph.add_run(name)

            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(name_size)

            paragraph.paragraph_format.space_after = Pt(2)

        else:

            # Classic ATS

            paragraph = document.add_paragraph()

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            run = paragraph.add_run(name)

            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(name_size)

            paragraph.paragraph_format.space_after = Pt(3)

    # =====================================================
    # PROCESS CONTENT
    # =====================================================

    for raw_line in lines[1:]:

        line = raw_line.strip()

        if not line:
            continue

        # -------------------------------------------------
        # CONTACT INFORMATION
        # -------------------------------------------------

        if (
            "@" in line
            or "linkedin.com" in line.lower()
            or "github.com" in line.lower()
            or line.startswith("+91")
        ):

            paragraph = document.add_paragraph()

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            run = paragraph.add_run(line)

            run.font.name = font_name
            run.font.size = Pt(9)

            if template == "Modern Professional":

                run.font.color.rgb = accent_color

            paragraph.paragraph_format.space_after = Pt(5)

            continue

        # -------------------------------------------------
        # SECTION HEADINGS
        # -------------------------------------------------

        if line.upper() in SECTION_NAMES:

            paragraph = document.add_paragraph()

            if template == "Modern Professional":

                paragraph.paragraph_format.space_before = Pt(9)
                paragraph.paragraph_format.space_after = Pt(3)

            elif template == "Minimal Student":

                paragraph.paragraph_format.space_before = Pt(5)
                paragraph.paragraph_format.space_after = Pt(2)

            else:

                paragraph.paragraph_format.space_before = Pt(8)
                paragraph.paragraph_format.space_after = Pt(3)

            run = paragraph.add_run(
                line.upper()
            )

            run.bold = True
            run.font.name = font_name
            run.font.size = Pt(heading_size)
            run.font.color.rgb = accent_color

            if template == "Modern Professional":

                add_bottom_border(
                    paragraph,
                    color="1F4E79",
                    size="8"
                )

            elif template == "Minimal Student":

                add_bottom_border(
                    paragraph,
                    color="000000",
                    size="4"
                )

            else:

                add_bottom_border(
                    paragraph,
                    color="000000",
                    size="6"
                )

            continue

        # -------------------------------------------------
        # BULLET POINTS
        # -------------------------------------------------

        if line.startswith("-") or line.startswith("•"):

            clean_line = line.lstrip("-•").strip()

            paragraph = document.add_paragraph(
                style="List Bullet"
            )

            if template == "Minimal Student":

                paragraph.paragraph_format.space_after = Pt(1)

            else:

                paragraph.paragraph_format.space_after = Pt(2)

            run = paragraph.add_run(
                clean_line
            )

            run.font.name = font_name
            run.font.size = Pt(
                9 if template == "Minimal Student"
                else 10
            )

            continue

        # -------------------------------------------------
        # NORMAL CONTENT
        # -------------------------------------------------

        paragraph = document.add_paragraph()

        if template == "Minimal Student":

            paragraph.paragraph_format.space_after = Pt(2)

        else:

            paragraph.paragraph_format.space_after = Pt(3)

        run = paragraph.add_run(line)

        run.font.name = font_name
        run.font.size = Pt(body_size)

    # =====================================================
    # SAVE
    # =====================================================

    file_path = "AI_Resume.docx"

    document.save(file_path)

    return file_path